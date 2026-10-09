import io
import uuid
import openpyxl
import pytest
from app.report.excel_generator import build_chain_sheet_entries, generate_chain_sheet_excel
from app.report.report_builder import ReportBuilder


def test_generate_chain_sheet_excel_structure_and_styling():
    report_data = {
        "run_id": "run-test-1508692",
        "state": "FL",
        "county": "miami-dade",
        "query_type": "address",
        "query_value": "620 arvida pkwy",
        "property": {
            "apn": "03-5105-002-0230",
            "address": "620 arvida pkwy, coral gables, FL 33156",
            "owner_name": "JOHN H RUIZ",
            "legal_description": "GABLES ESTS NO 3 PB 65-66 LOT 25 BLK C",
        },
        "tax_record": {
            "status": "Paid",
            "gross_tax": 18450.25,
            "tax_account": "03-5105-002-0230",
        },
        "recorder_primary_documents": [
            {
                "grantor": "AP 620 LLC",
                "grantee": "BENJAMIN LEON JR AND SILVIA LEON",
                "document_type": "WARRANTY DEED",
                "recording_date": "2010-12-21",
                "book_page": "27529/1785",
                "instrument_number": "2010 R 853064",
                "comments": "Lot 25",
                "ocr_json": {"source": "recorder"},
            },
            {
                "grantor": "BENJAMIN LEON JR AND SILVIA LEON",
                "grantee": "JOHN H RUIZ",
                "document_type": "WARRANTY DEED",
                "recording_date": "2020-04-13",
                "book_page": "31894/4474",
                "instrument_number": "2020 R 223930",
                "comments": "Lot : 24,25",
                "ocr_json": {"source": "recorder"},
            },
            {
                "grantor": "JOHN H RUIZ",
                "grantee": "CITY NATIONAL BANK OF FLORIDA",
                "document_type": "MORTGAGE",
                "recording_date": "2022-04-12",
                "book_page": "33126/1399",
                "instrument_number": "2022 R 301183",
                "amount": 20000000.0,
                "comments": "$20,000,000.00",
                "ocr_json": {"source": "recorder"},
            },
        ],
    }

    excel_io = generate_chain_sheet_excel(report_data)
    assert isinstance(excel_io, io.BytesIO)
    assert excel_io.getbuffer().nbytes > 0

    # Load workbook back to inspect structure
    wb = openpyxl.load_workbook(excel_io)
    assert "Sheet1" in wb.sheetnames
    assert "Command" in wb.sheetnames

    ws1 = wb["Sheet1"]

    # Check top metadata
    assert ws1["A1"].value == "Search Date:"
    assert ws1["A4"].value == "Tax Information"
    assert ws1["B4"].value == "ONLINE"
    assert ws1["A5"].value == "Account Identifier:"
    assert ws1["B5"].value == "03-5105-002-0230"
    assert ws1["A11"].value == "Owner Name:"
    assert "JOHN H RUIZ" in str(ws1["B11"].value)
    assert ws1["A13"].value == "County:"
    assert "Miami-Dade" in str(ws1["B13"].value)

    # Check table headers on Row 15
    assert ws1["A15"].value == "Grantor Name"
    assert ws1["B15"].value == "Grantee Name"
    assert ws1["C15"].value == "Order Type"
    assert ws1["D15"].value == "Record Date"
    assert ws1["E15"].value == "Book/Page"
    assert ws1["F15"].value == "Instrument#"
    assert ws1["G15"].value == "Comments"

    # Check peach fill on Order Type header
    header_fill = ws1["C15"].fill.start_color.rgb
    assert "FCE4D6" in str(header_fill).upper()

    # Rows are sorted newest first
    assert ws1["A16"].value == "JOHN H RUIZ"
    assert ws1["B16"].value == "CITY NATIONAL BANK OF FLORIDA"
    assert ws1["C16"].value == "MORTGAGE"
    assert ws1["D16"].value == "12-04-2022"
    assert ws1["E16"].value == "33126/1399"
    assert ws1["F16"].value == "2022 R 301183"

    assert ws1["A17"].value == "BENJAMIN LEON JR AND SILVIA LEON"
    assert ws1["B17"].value == "JOHN H RUIZ"
    assert ws1["C17"].value == "WARRANTY DEED"

    assert ws1["A18"].value == "AP 620 LLC"
    assert ws1["B18"].value == "BENJAMIN LEON JR AND SILVIA LEON"
    assert ws1["C18"].value == "WARRANTY DEED"
    assert ws1["D18"].value == "21-12-2010"
    assert ws1["E18"].value == "27529/1785"
    assert ws1["F18"].value == "2010 R 853064"
    assert ws1["G18"].value == "Lot 25"

    # Check Order Type cell fill on row 16
    row_fill = ws1["C16"].fill.start_color.rgb
    assert "FCE4D6" in str(row_fill).upper()

    # Check Mortgage amount comment on newest row
    assert "$20,000,000.00" in str(ws1["G16"].value)

    # Check Command sheet
    ws2 = wb["Command"]
    assert "CHAIN SHEET AUDIT" in str(ws2["A1"].value)
    assert ws2["B3"].value == "run-test-1508692"


def test_build_chain_sheet_entries_includes_name_search_documents():
    report_data = {
        "recorder_primary_documents": [
            {
                "document_type": "DEED - DEE",
                "recording_date": "2022-01-24",
                "book_page": "32977/324",
                "instrument_number": "2022 R 66756",
                "grantor": "LIFT STATIONS OF SOUTH FLORIDA LLC",
                "grantee": "MIAMI EDGE INVESTMENTS INC",
                "ocr_json": {"source": "recorder"},
            }
        ],
        "name_search_groups": [
            {
                "name": "MIAMI EDGE INVESTMENTS INC",
                "documents": [
                    {
                        "document_type": "MORTGAGE - MOR",
                        "recording_date": "2022-01-24",
                        "book_page": "32977/326",
                        "instrument_number": "2022 R 66757",
                        "grantor": "MIAMI EDGE INVESTMENTS INC",
                        "grantee": "RBI MORTGAGES LLC",
                        "ocr_json": {"source": "name_searcher", "searched_name": "MIAMI EDGE INVESTMENTS INC"},
                    }
                ],
            }
        ],
    }

    entries = build_chain_sheet_entries(report_data)
    book_pages = {_normalize_book_page(e.get("book_page")) for e in entries}
    assert "32977/324" in book_pages
    assert "32977/326" in book_pages


def test_generate_chain_sheet_excel_leaves_comments_empty_without_amount_or_notes():
    report_data = {
        "property": {
            "legal_description": "GARDEN CITY LOT 24, 25 BLK 9",
        },
        "recorder_primary_documents": [
            {
                "document_type": "DEED - DEE",
                "recording_date": "2022-01-24",
                "book_page": "32977/324",
                "instrument_number": "2022 R 66756",
                "grantor": "LIFT STATIONS OF SOUTH FLORIDA LLC",
                "grantee": "MIAMI EDGE INVESTMENTS INC",
                "ocr_json": {"source": "recorder"},
            }
        ],
    }

    excel_io = generate_chain_sheet_excel(report_data)
    wb = openpyxl.load_workbook(excel_io)
    ws1 = wb["Sheet1"]

    assert ws1["G16"].value in (None, "")


def test_build_chain_sheet_entries_excludes_assessor_only_sales():
    report_data = {
        "chain_of_title": [
            {
                "document_type": "Sale",
                "recording_date": "2017-06-13",
                "book_page": "30576-2086",
                "sale_price": 4268000,
            },
            {
                "document_type": "Sale",
                "recording_date": "1971-01-01",
                "book_page": "18859-3923",
                "sale_price": 16000,
            },
        ],
        "recorder_primary_documents": [],
        "name_search_groups": [],
    }

    entries = build_chain_sheet_entries(report_data)
    assert entries == []


def test_build_chain_sheet_entries_enriches_assessor_sales_with_recorder_deed():
    report_data = {
        "chain_of_title": [
            {
                "document_type": "Sale",
                "recording_date": "2019-10-30",
                "book_page": "31687-1679",
                "sale_price": 387500,
            },
            {
                "document_type": "Sale",
                "recording_date": "2017-06-13",
                "book_page": "30576-2086",
                "sale_price": 4268000,
            },
        ],
        "recorder_primary_documents": [
            {
                "document_type": "DEED - DEE",
                "recording_date": "2019-11-13",
                "book_page": "31687/1679",
                "instrument_number": "2019 R 711430",
                "grantor": "D R HORTON INC",
                "grantee": "MORALES JUAN A",
                "ocr_json": {"source": "recorder"},
            }
        ],
    }

    entries = build_chain_sheet_entries(report_data)
    assert len(entries) == 1

    vesting = next(e for e in entries if _normalize_book_page(e.get("book_page")) == "31687/1679")
    assert vesting["grantor"] == "D R HORTON INC"
    assert vesting["grantee"] == "MORALES JUAN A"
    assert vesting["instrument_number"] == "2019 R 711430"
    assert vesting["document_type"] == "DEED - DEE"
    assert vesting["sale_price"] == 387500


def _normalize_book_page(value):
    from app.report.excel_generator import _normalize_book_page_key

    return _normalize_book_page_key(value)


def test_generate_chain_sheet_excel_includes_recorder_grantor_grantee():
    report_data = {
        "run_id": "run-test-chain-merge",
        "property": {
            "apn": "30-6924-001-0030",
            "owner_name": "JUAN A MORALES",
            "legal_description": "SUMMERVILLE VILLAS LOT 3 BLK 1",
        },
        "chain_of_title": [
            {
                "document_type": "Sale",
                "recording_date": "2019-10-30",
                "book_page": "31687-1679",
                "sale_price": 387500,
            }
        ],
        "recorder_primary_documents": [
            {
                "document_type": "DEED - DEE",
                "recording_date": "2019-11-13",
                "book_page": "31687/1679",
                "instrument_number": "2019 R 711430",
                "grantor": "D R HORTON INC",
                "grantee": "MORALES JUAN A",
            }
        ],
        "name_searches": [
            {"name": "MORALES JUAN A", "source": "31687/1679"},
            {"name": "D R HORTON INC", "source": "31687/1679"},
        ],
    }

    excel_io = generate_chain_sheet_excel(report_data)
    wb = openpyxl.load_workbook(excel_io)
    ws1 = wb["Sheet1"]

    assert ws1["A16"].value == "D R HORTON INC"
    assert ws1["B16"].value == "MORALES JUAN A"
    assert ws1["C16"].value == "DEED - DEE"
    assert ws1["F16"].value == "2019 R 711430"
    assert "$387,500.00" in str(ws1["G16"].value)


@pytest.mark.asyncio
async def test_reports_excel_download_route():
    from fastapi.testclient import TestClient
    from app.api.main import app
    from app.db.supabase_client import get_memory_store

    run_id = f"test-excel-run-{uuid.uuid4().hex[:8]}"
    run_dict = {
        "id": run_id,
        "state": "FL",
        "county": "miami-dade",
        "query_type": "address",
        "query_value": "620 arvida pkwy",
        "parcel": "03-5105-002-0230",
        "address": "620 arvida pkwy",
        "owner_name": "JOHN H RUIZ",
        "status": "completed",
        "plan_json": {
            "node_results": {
                "assessor": {
                    "property": {
                        "apn": "03-5105-002-0230",
                        "address": "620 arvida pkwy",
                        "owner_name": "JOHN H RUIZ",
                        "legal_description": "GABLES ESTS NO 3 PB 65-66 LOT 25 BLK C",
                    }
                }
            }
        },
    }

    get_memory_store().runs[run_id] = run_dict

    client = TestClient(app)
    response = client.get(f"/reports/run/{run_id}/excel/download")

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers.get("content-type", "")
    assert "attachment" in response.headers.get("content-disposition", "")
    assert ".xlsx" in response.headers.get("content-disposition", "")

    # Ensure returned stream is a valid workbook
    content = io.BytesIO(response.content)
    wb = openpyxl.load_workbook(content)
    assert "Sheet1" in wb.sheetnames
    assert "Command" in wb.sheetnames
